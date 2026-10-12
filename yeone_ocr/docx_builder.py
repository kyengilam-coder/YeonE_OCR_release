"""python-docx 로 DOCX 도 함께 만들 수 있다(선택). HWPX 와 같은 꾸밈: 쪽 나누기, 확인 등급 색(초록·파랑·빨강)과 밑줄,
고지 사항은 Word 주석, 작업 표시 회색."""
from __future__ import annotations
from pathlib import Path
from .blocks import Block
from .hwpx_builder import figure_width_mm, split_marks, FONT_BODY, FONT_HEAD, MEMO_AUTHOR
from .grades import styled_runs, GRADE_STYLE, STYLE_ALIAS

_UNDERLINE = {"DOT": "DOTTED", "DASH": "DASH", "DOUBLE_SLIM": "DOUBLE", "SOLID": "SINGLE", "DASH_DOT": "DOT_DASH", "WAVE": "WAVY"}


def _grade(run, g: dict) -> None:
    """등급 모양: 글자색·굵게·밑줄, 형광펜은 노랑이면 형광펜, 그 밖의 색은 음영(w:shd)."""
    from docx.shared import RGBColor
    from docx.enum.text import WD_COLOR_INDEX, WD_UNDERLINE
    from docx.oxml.ns import qn
    run.font.color.rgb = RGBColor.from_string(g["color"].lstrip("#")); run.font.bold = g["bold"]
    if g["underline"]:
        run.font.underline = getattr(WD_UNDERLINE, _UNDERLINE.get(g["underline"], "SINGLE"))
    if g["highlight"] and g["highlight"].upper() == "#FFFF00":
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    elif g["highlight"]:
        rpr = run._element.get_or_add_rPr()
        shd = rpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): g["highlight"].lstrip("#")}); rpr.append(shd)


def _font(run, name: str, size_pt: float, bold=False, color: str | None = None, highlight=False):
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml.ns import qn
    run.font.name = name; run.font.size = Pt(size_pt); run.font.bold = bold
    rpr = run._element.get_or_add_rPr(); rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {}); rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def build_docx(blocks: list[Block], out_path: str, min_mm: float = 70, large_min_mm: float = 120,
               highlight: bool = True, page_breaks: bool = True, memos: bool = True) -> dict:
    import docx
    from docx.shared import Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    D = docx.Document()
    for s in D.sections:
        s.left_margin = s.right_margin = Mm(25); s.top_margin = Mm(22); s.bottom_margin = Mm(20)
    stats = {"paragraphs": 0, "headings": 0, "images": 0, "tables": 0, "marks": 0, "page_breaks": 0, "memos": 0}
    can_comment = memos and hasattr(D, "add_comment")             # python-docx 1.2 부터

    def para(text, size, *, bold=False, color=None, marks=False, font=FONT_BODY, spans=(), notes=()):
        p = D.add_paragraph(); stats["paragraphs"] += 1
        if notes and not can_comment:                              # 주석을 못 쓰면 고지를 글 끝에 회색으로
            text = text + "".join(f"〔고지: {m}〕" for _, m in notes); notes = ()
        if marks or spans or notes:
            stats["marks"] += sum(1 for _, k in split_marks(text) if k == "mark")
            for seg, kind, msgs in styled_runs(text, spans if highlight else (), notes):
                kind = STYLE_ALIAS.get(kind, kind)
                r = p.add_run(seg)
                if kind == "note":
                    _font(r, font, 9, color="#666666")
                elif kind == "body" or not highlight:
                    _font(r, font, size, bold=bold, color=color)
                else:
                    _font(r, font, size); _grade(r, GRADE_STYLE[kind])
                for m in msgs:
                    D.add_comment(r, text=m, author=MEMO_AUTHOR, initials="연"); stats["memos"] += 1
        else:
            _font(p.add_run(text), font, size, bold=bold, color=color)
        return p

    first_meta = True
    for b in blocks:
        if b.type == "meta":
            if page_breaks and not first_meta:
                D.add_page_break(); stats["page_breaks"] += 1
            first_meta = False
            para(b.text, 8, color="#8A8A8A", font=FONT_HEAD, notes=b.memos)
        elif b.type == "heading":
            size = 16 if b.level <= 1 else (14 if b.level == 2 else 12)
            para(b.text, size, bold=True, font=FONT_HEAD); stats["headings"] += 1
        elif b.type == "para":
            para(b.text, 9.5 if b.kind == "소주" else 11, marks=True, spans=b.spans, notes=b.memos)
        elif b.type == "note":
            para(b.text, 9, color="#666666", marks=True)
        elif b.type == "image":
            w = figure_width_mm(b, min_mm, large_min_mm)
            if w * b.ratio > 230: w = 230 / b.ratio
            D.add_picture(b.path, width=Mm(w)); D.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; stats["images"] += 1
        elif b.type == "table":
            t = D.add_table(rows=len(b.rows), cols=len(b.rows[0])); t.style = "Table Grid"
            for i, row in enumerate(b.rows):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell
            stats["tables"] += 1
        elif b.type == "uncertain":
            para(b.text or "[판독 확인 필요]", 9.5, bold=True, color="#666666", font=FONT_HEAD)
            for it in b.items:
                para("· " + it, 9, color="#666666", marks=True)
        elif b.type == "report":
            para("〔작업 기록〕", 14, bold=True, font=FONT_HEAD)
            for line in b.text.splitlines():
                para(line, 9, color="#666666")
    D.save(out_path)
    return stats
